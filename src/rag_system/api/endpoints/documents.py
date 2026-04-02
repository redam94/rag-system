"""Document upload and listing endpoints."""

import io
import uuid
from typing import List

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from loguru import logger

from ..dependencies import get_rag, get_session_id
from ..models import DocumentListItem, DocumentListResponse, DocumentUploadResponse
from ...memory.context_rag import ContextRAG

router = APIRouter()

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".pptx"}


def _extract_text(filename: str, content: bytes) -> tuple[str, str]:
    """Extract text from uploaded file bytes. Returns (text, source_type)."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in (".txt", ".md"):
        return content.decode("utf-8", errors="replace"), ext.lstrip(".")

    if ext == ".pdf":
        try:
            import fitz  # pymupdf
        except ImportError:
            raise HTTPException(500, "pymupdf is required for PDF uploads")
        doc = fitz.open(stream=content, filetype="pdf")
        text = "\n\n".join(page.get_text() for page in doc)
        doc.close()
        return text, "pdf"

    if ext == ".pptx":
        try:
            from pptx import Presentation
        except ImportError:
            raise HTTPException(500, "python-pptx is required for PPTX uploads")
        prs = Presentation(io.BytesIO(content))
        parts: List[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return "\n\n".join(parts), "pptx"

    raise HTTPException(
        415,
        f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}",
    )


@router.post("/upload", response_model=DocumentUploadResponse)
async def upload_document(
    file: UploadFile = File(...),
    session_id: str = Depends(get_session_id),
):
    """Upload a document to the RAG store."""
    rag = get_rag()

    if not file.filename:
        raise HTTPException(400, "Filename is required")

    raw = await file.read()
    if not raw:
        raise HTTPException(400, "Empty file")

    text, source_type = _extract_text(file.filename, raw)
    if not text.strip():
        raise HTTPException(400, "No text could be extracted from the file")

    workflow_id = f"upload_{uuid.uuid4().hex[:8]}"

    # ContextRAG.add_document returns None but adds chunks internally
    rag.add_document(
        content=text,
        title=file.filename,
        source_type=source_type,
        workflow_id=workflow_id,
        source_path=file.filename,
    )

    # Estimate chunk count from the chunker
    chunks = rag.chunker.chunk_text(text, {"title": file.filename})
    chunk_count = len(chunks)

    logger.info(f"Uploaded document '{file.filename}' ({source_type}, {chunk_count} chunks)")

    return DocumentUploadResponse(
        filename=file.filename,
        source_type=source_type,
        chunks_added=chunk_count,
        message=f"Successfully uploaded '{file.filename}'",
    )


@router.get("", response_model=DocumentListResponse)
async def list_documents(session_id: str = Depends(get_session_id)):
    """List all documents in the RAG store."""
    rag = get_rag()

    if not rag.enabled:
        return DocumentListResponse(documents=[], total=0)

    try:
        # Query all document-type entries from ChromaDB
        results = rag.collection.get(
            where={"type": "document"},
            include=["metadatas"],
        )
    except Exception:
        # If no documents with type filter, return empty
        return DocumentListResponse(documents=[], total=0)

    # Aggregate by title
    title_counts: dict[str, dict] = {}
    for meta in results.get("metadatas", []):
        title = meta.get("title", "unknown")
        if title not in title_counts:
            title_counts[title] = {
                "title": title,
                "source_type": meta.get("source_type", "unknown"),
                "chunk_count": 0,
            }
        title_counts[title]["chunk_count"] += 1

    documents = [DocumentListItem(**v) for v in title_counts.values()]
    return DocumentListResponse(documents=documents, total=len(documents))
